#!/usr/bin/env python3
"""
Add a time-vs-weight scatter chart (points only, no connecting line) to
every *_weights.xlsx under batch_output/. Columns are frame_no,
timestamp_sec, raw_digits, weight (written by extract_lcd_weights_batch.py).
Run after processing videos. Safe to re-run - replaces any chart already on
the sheet instead of stacking duplicates.
"""
import sys
sys.path.append('/home/bee/projeler/lcd_reader/.venv/lib/python3.14/site-packages')

import glob
from pathlib import Path
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties

MARKER_COLOR = "1F77B4"

for path in sorted(glob.glob(str(Path(__file__).parent / "batch_output" / "*" / "*_weights.xlsx"))):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    n_rows = ws.max_row

    ws._charts = []  # drop any chart already placed by a previous run

    chart = ScatterChart()
    chart.title = f"{Path(path).stem} - weight vs time"
    chart.x_axis.title = "time (s)"
    chart.y_axis.title = "weight (g)"
    # ScatterChart doesn't auto-assign axis positions like LineChart does -
    # both default to "l" otherwise, which renders as a squashed, near-empty
    # plot in LibreOffice/Excel.
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.width = 24
    chart.height = 10
    chart.style = 2

    x_ref = Reference(ws, min_col=2, min_row=2, max_row=n_rows)  # B: timestamp_sec
    y_ref = Reference(ws, min_col=4, min_row=1, max_row=n_rows)  # D: weight (incl. header, becomes series title)
    series = Series(y_ref, x_ref, title_from_data=True)
    series.marker = Marker(symbol="circle", size=5)
    # A marker with no explicit fill renders transparent (invisible) in
    # LibreOffice/Excel - has to be given a color explicitly.
    series.marker.graphicalProperties = GraphicalProperties(solidFill=MARKER_COLOR)
    series.marker.graphicalProperties.line.solidFill = MARKER_COLOR
    series.graphicalProperties.line.noFill = True  # points only, no connecting line

    chart.series.append(series)
    ws.add_chart(chart, "F2")
    wb.save(path)
    print(f"chart added: {path}")
