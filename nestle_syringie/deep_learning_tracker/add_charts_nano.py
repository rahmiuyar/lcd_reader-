#!/usr/bin/env python3
"""
Add a time-vs-weight scatter chart (points only, no connecting line) to
every *_weights.xlsx under a given output folder (default: test_nano_full/).
Same chart logic as ../add_charts.py, just pointed at whichever of this
experiment's output folders is passed in - each nano_tracker script version
(test_nano_full/, test_nano_full_v2/, test_nano_full_v3/, ...) gets its own
folder, so this takes the target as an argument instead of hardcoding one.
Safe to re-run - replaces any chart already on the sheet instead of
stacking duplicates.

Usage: ../../.venv/bin/python add_charts_nano.py [folder_name]
  (folder_name is relative to this script's directory, default test_nano_full)
"""
import sys
sys.path.append('/home/bee/projeler/lcd_reader/.venv/lib/python3.14/site-packages')

import glob
from pathlib import Path
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties

MARKER_COLOR = "1F77B4"

target_dir = sys.argv[1] if len(sys.argv) > 1 else "test_nano_full"

for path in sorted(glob.glob(str(Path(__file__).parent / target_dir / "*" / "*_weights.xlsx"))):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    n_rows = ws.max_row

    ws._charts = []  # drop any chart already placed by a previous run

    chart = ScatterChart()
    chart.title = f"{Path(path).stem} - weight vs time"
    chart.x_axis.title = "time (s)"
    chart.y_axis.title = "weight (g)"
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.width = 24
    chart.height = 10
    chart.style = 2

    x_ref = Reference(ws, min_col=2, min_row=2, max_row=n_rows)  # B: timestamp_sec
    y_ref = Reference(ws, min_col=4, min_row=1, max_row=n_rows)  # D: weight (incl. header, becomes series title)
    series = Series(y_ref, x_ref, title_from_data=True)
    series.marker = Marker(symbol="circle", size=5)
    series.marker.graphicalProperties = GraphicalProperties(solidFill=MARKER_COLOR)
    series.marker.graphicalProperties.line.solidFill = MARKER_COLOR
    series.graphicalProperties.line.noFill = True  # points only, no connecting line

    chart.series.append(series)
    ws.add_chart(chart, "F2")
    wb.save(path)
    print(f"chart added: {path}")
